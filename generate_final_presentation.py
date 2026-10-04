import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = pptx.Presentation()
    # 16:9 Widescreen standard dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Colors
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A Primary dark header
    BLUE = RGBColor(37, 99, 235)       # #2563EB Accent blue
    DARK = RGBColor(15, 23, 42)        # #0F172A Body text
    MUTED = RGBColor(71, 85, 105)      # #475569 Subtitles
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC Card background
    BORDER = RGBColor(226, 232, 240)   # #E2E8F0 Card borders
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(22, 163, 74)      # #16A34A Success
    RED = RGBColor(220, 38, 38)        # #DC2626 Warning
    
    def add_header(slide, title, subtitle):
        # Header text box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = NAVY
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = 'Arial'
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(3)

    def add_card(slide, left, top, width, height, fill_color=LIGHT_BG, border_color=BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    
    # Background accent bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.4))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY
    top_bar.line.fill.background()
    
    # Title card container
    add_card(s1, 1.0, 1.0, 11.333, 5.5, fill_color=LIGHT_BG, border_color=BORDER)
    
    # Category / Course Banner
    tb_c = s1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.333), Inches(0.5))
    tf_c = tb_c.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.text = "MINI PROJECT  |  23CS4219 - SOFTWARE ENGINEERING LABORATORY"
    p_c.font.name = 'Arial'
    p_c.font.size = Pt(12)
    p_c.font.bold = True
    p_c.font.color.rgb = BLUE
    p_c.alignment = PP_ALIGN.CENTER
    
    # Project Title
    tb_t = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(1.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "Campus Complaint & Maintenance Management System"
    p_t.font.name = 'Arial'
    p_t.font.size = Pt(30)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY
    p_t.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    p_sub = tf_t.add_paragraph()
    p_sub.text = "A Role-Based Full-Stack Web Platform for Institutional Grievance Lifecycle Management"
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = MUTED
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(6)
    
    # Metadata card grid
    meta_box = s1.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(10.333), Inches(2.6))
    tf_m = meta_box.text_frame
    tf_m.word_wrap = True
    
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "Submitted by:  KARTHIKEYAN GULLIPALLI ([ROLL NUMBER])"
    p_m1.font.name = 'Arial'
    p_m1.font.size = Pt(14)
    p_m1.font.bold = True
    p_m1.font.color.rgb = DARK
    p_m1.alignment = PP_ALIGN.CENTER
    
    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Under the Guidance of:  [GUIDE NAME]"
    p_m2.font.name = 'Arial'
    p_m2.font.size = Pt(13)
    p_m2.font.color.rgb = DARK
    p_m2.alignment = PP_ALIGN.CENTER
    p_m2.space_before = Pt(4)
    
    p_m3 = tf_m.add_paragraph()
    p_m3.text = "Department of Computer Science and Engineering"
    p_m3.font.name = 'Arial'
    p_m3.font.size = Pt(13)
    p_m3.font.bold = True
    p_m3.font.color.rgb = NAVY
    p_m3.alignment = PP_ALIGN.CENTER
    p_m3.space_before = Pt(12)
    
    p_m4 = tf_m.add_paragraph()
    p_m4.text = "Anil Neerukonda Institute of Technology and Sciences (ANITS), Visakhapatnam\nAcademic Year: 2024–2025"
    p_m4.font.name = 'Arial'
    p_m4.font.size = Pt(11.5)
    p_m4.font.color.rgb = MUTED
    p_m4.alignment = PP_ALIGN.CENTER
    p_m4.space_before = Pt(3)

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & OBJECTIVES
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement & Project Objectives", "Addressing institutional operational bottlenecks through structured digitization")
    
    # Left Card: Problem Statement
    add_card(s2, 0.8, 1.7, 5.7, 5.2)
    tb_p = s2.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.8))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    
    p = tf_p.paragraphs[0]
    p.text = "The Core Problems (Manual Approach)"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = RED
    
    bullets_p = [
        ("Physical Vulnerability: ", "Paper registers are vulnerable to damage, misplacement, and page loss."),
        ("Zero Visibility: ", "Complainants have no tracking mechanism to monitor repair progress."),
        ("Lack of Accountability: ", "Technicians receive verbal instructions without timestamps or audit trails."),
        ("Ambiguous Fault Locations: ", "Handwritten complaints lack floor/room context, causing wasted repair efforts."),
        ("Infeasible Historical Auditing: ", "Generating response times or recurring equipment failure reports is nearly impossible.")
    ]
    for b_prefix, b_text in bullets_p:
        p_b = tf_p.add_paragraph()
        p_b.space_before = Pt(10)
        r1 = p_b.add_run()
        r1.text = "• " + b_prefix; r1.font.bold = True; r1.font.size = Pt(12); r1.font.color.rgb = DARK
        r2 = p_b.add_run()
        r2.text = b_text; r2.font.size = Pt(12); r2.font.color.rgb = MUTED

    # Right Card: Objectives
    add_card(s2, 6.8, 1.7, 5.7, 5.2)
    tb_o = s2.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.8))
    tf_o = tb_o.text_frame
    tf_o.word_wrap = True
    
    p = tf_o.paragraphs[0]
    p.text = "Key Project Objectives"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BLUE
    
    bullets_o = [
        ("Centralized Digitization: ", "Eliminate physical logbooks by offering a 24/7 web-based grievance portal."),
        ("Strict Role-Based Access: ", "Implement RBAC supporting Student, Faculty, Maintenance, and Admin."),
        ("Finite State Lifecycle: ", "Enforce audited state progression: NEW → VERIFIED → ASSIGNED → IN_PROGRESS → RESOLVED → CLOSED."),
        ("Location & Urgency Mapping: ", "Categorize tickets by Building, Floor, Room, and Priority (Low to Critical)."),
        ("Technician Task Queue: ", "Equip maintenance staff with active task queues and repair remarks logging."),
        ("Administrative Governance: ", "Provide real-time analytics across category breakdowns and open hazards.")
    ]
    for b_prefix, b_text in bullets_o:
        p_b = tf_o.add_paragraph()
        p_b.space_before = Pt(8)
        r1 = p_b.add_run()
        r1.text = "✓ " + b_prefix; r1.font.bold = True; r1.font.size = Pt(12); r1.font.color.rgb = DARK
        r2 = p_b.add_run()
        r2.text = b_text; r2.font.size = Pt(12); r2.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 3: EXISTING SYSTEM vs PROPOSED SYSTEM
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Existing System vs. Proposed System", "Comparative evaluation of operational efficiency and data transparency")
    
    # Table layout
    rows = 7; cols = 3
    t_shape = s3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.1))
    tbl = t_shape.table
    tbl.columns[0].width = Inches(2.3)
    tbl.columns[1].width = Inches(4.7)
    tbl.columns[2].width = Inches(4.733)
    
    comp_headers = ["Operational Dimension", "Existing System (Manual Registers)", "Proposed System (CCMS Web Portal)"]
    for c_idx, h_text in enumerate(comp_headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if c_idx != 1 else RGBColor(185, 28, 28)
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.name = 'Arial'; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
    comp_rows = [
        ["Grievance Submission", "Physical visit to department counter; manual paper register entry during office hours only.", "24/7 web access; structured categorization, campus location mapping, and optional photo attachment."],
        ["Status Tracking", "No tracking mechanism; requires repeated verbal follow-ups at the caretaker desk.", "Real-time status badges, chronological audit log, and assigned technician details."],
        ["Verification & Triage", "Supervisors read through handwritten logbooks; no filtering or formal verification step.", "Interactive admin triage: instant search, status filtering, one-click verify or reject with remarks."],
        ["Task Dispatching", "Verbal work instructions; frequent ambiguity regarding exact room number or priority.", "Direct digital assignment to qualified maintenance technician with supervisor notes."],
        ["Work Execution", "No formal work orders; repair completion communicated verbally or not at all.", "Technician task queue; status update to IN_PROGRESS and diagnostic remarks upon resolution."],
        ["Satisfaction Review", "Complainants have no mechanism to evaluate repair quality or reopen issues.", "1 to 5 star rating submission, evaluative commentary, and reopening capability if issue persists."]
    ]
    
    for r_idx, row_vals in enumerate(comp_rows):
        for c_idx, val in enumerate(row_vals):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 0 else LIGHT_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = 'Arial'; p.font.size = Pt(10.5)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = NAVY
            elif c_idx == 1:
                p.font.color.rgb = RGBColor(127, 29, 29)
            else:
                p.font.color.rgb = RGBColor(20, 83, 45)

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "System Architecture", "Decoupled 3-tier client-server architecture with stateless JWT security")
    
    # Left container for image
    add_card(s4, 0.8, 1.7, 4.2, 5.2, fill_color=WHITE)
    if os.path.exists("docs/diagrams/fig04_architecture.png"):
        s4.shapes.add_picture("docs/diagrams/fig04_architecture.png", Inches(1.3), Inches(1.85), height=Inches(4.9))
        
    # Right container: Layer breakdown
    add_card(s4, 5.3, 1.7, 7.233, 5.2)
    tb_arch = s4.shapes.add_textbox(Inches(5.5), Inches(1.9), Inches(6.833), Inches(4.8))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True
    
    p = tf_arch.paragraphs[0]
    p.text = "Architectural Tier Breakdown"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = NAVY
    
    layers = [
        ("Presentation Tier (React 18 SPA)", "Single Page Application built with Vite 5 and styled via Tailwind CSS. Axios client automatically injects JWT Bearer token into all API headers for seamless authorization."),
        ("Application Tier (Node.js & Express)", "Stateless REST API implementing CORS, express-validator sanitization, Multer multipart upload handler, and verifyToken / requireRole middleware pipelines."),
        ("Storage Tier (MySQL 8 Relational DB)", "ACID-compliant relational database (ccms_db) organized into 7 normalized tables. Connected via promise-based connection pool executing parameterized queries."),
        ("Security & Session Model", "Stateless authentication via JSON Web Tokens (JWT) signed with HS256. Passwords securely hashed and salted using bcryptjs (10 rounds).")
    ]
    for l_title, l_desc in layers:
        p_l = tf_arch.add_paragraph()
        p_l.space_before = Pt(10)
        r1 = p_l.add_run(); r1.text = "• " + l_title + "\n"; r1.font.bold = True; r1.font.size = Pt(11.5); r1.font.color.rgb = BLUE
        r2 = p_l.add_run(); r2.text = l_desc; r2.font.size = Pt(10.5); r2.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 5: MAJOR MODULES
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Major System Modules", "Cohesive functional modules mapped to client routes and REST controllers")
    
    modules = [
        ("1. Authentication & RBAC Module", "Manages user registration, bcryptjs password hashing, JWT token generation, and role authorization guards across Student, Faculty, Maintenance, and Admin."),
        ("2. Complaint Lodging Module", "Provides reactive form controls mapping issues to standardized campus categories and physical rooms, supporting Multer image uploads and generating unique CMP-IDs."),
        ("3. Admin Triage & Dashboard", "Equips administrators with campus-wide ticket visibility, multi-filter search, one-click verification/rejection with remarks, and live metric summaries."),
        ("4. Maintenance Task Queue", "Dispatches verified tickets to active maintenance staff, providing technicians with focused task queues to start work and log resolution remarks."),
        ("5. Status Auditing & Feedback", "Maintains an immutable append-only audit log in complaint_updates for every transition, and collects 1–5 star user feedback ratings upon ticket resolution."),
        ("6. Master Data Administration", "Allows administrators to catalog campus buildings, floors, and rooms, manage active complaint categories, and toggle user active states without deleting history.")
    ]
    
    positions = [
        (0.8, 1.7), (4.8, 1.7), (8.8, 1.7),
        (0.8, 4.4), (4.8, 4.4), (8.8, 4.4)
    ]
    
    for idx, (m_title, m_desc) in enumerate(modules):
        x, y = positions[idx]
        add_card(s5, x, y, 3.733, 2.5)
        tb = s5.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(3.333), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.name = 'Arial'; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = NAVY
        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.name = 'Arial'; p2.font.size = Pt(10.5); p2.font.color.rgb = MUTED
        p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 6: COMPLAINT WORKFLOW / DFD
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Complaint Workflow & State Lifecycle", "Audited Finite State Machine enforcing valid operational progression")
    
    add_card(s6, 0.8, 1.7, 6.8, 5.2, fill_color=WHITE)
    if os.path.exists("docs/diagrams/fig11_state.png"):
        s6.shapes.add_picture("docs/diagrams/fig11_state.png", Inches(1.1), Inches(2.5), width=Inches(6.2))
        
    add_card(s6, 7.8, 1.7, 4.733, 5.2)
    tb_wf = s6.shapes.add_textbox(Inches(8.0), Inches(1.9), Inches(4.333), Inches(4.8))
    tf_wf = tb_wf.text_frame
    tf_wf.word_wrap = True
    
    p = tf_wf.paragraphs[0]
    p.text = "Lifecycle State Transitions"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = NAVY
    
    fsm_steps = [
        ("1. NEW: ", "Complaint lodged by student/faculty with location and description."),
        ("2. VERIFIED: ", "Administrator reviews details and authenticates the issue."),
        ("3. REJECTED: ", "Administrator dismisses invalid/out-of-scope ticket with notes."),
        ("4. ASSIGNED: ", "Administrator allocates verified ticket to designated technician."),
        ("5. IN_PROGRESS: ", "Technician accepts ticket and initiates physical remediation."),
        ("6. RESOLVED: ", "Technician completes work and logs diagnostic remarks."),
        ("7. CLOSED: ", "Complainant or Admin confirms fix; student prompted for rating."),
        ("8. REOPENED: ", "User reports issue still persists; ticket returns for reassignment.")
    ]
    for s_title, s_desc in fsm_steps:
        p_s = tf_wf.add_paragraph()
        p_s.space_before = Pt(5)
        r1 = p_s.add_run(); r1.text = s_title; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = BLUE
        r2 = p_s.add_run(); r2.text = s_desc; r2.font.size = Pt(10.5); r2.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 7: TECHNOLOGY STACK
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Technology Stack", "Robust, open-source technologies utilized in the CCMS implementation")
    
    stacks = [
        ("Frontend Technologies", [
            ("React 18.2: ", "Single Page Application component framework"),
            ("Vite 5.0: ", "High-performance build tooling & development server"),
            ("Tailwind CSS 3.3: ", "Utility-first responsive styling system"),
            ("React Router v6: ", "Client-side route navigation & role guards"),
            ("Axios 1.6: ", "HTTP client with JWT request interceptor")
        ]),
        ("Backend Technologies", [
            ("Node.js 18+ LTS: ", "Event-driven asynchronous server runtime"),
            ("Express.js 4.19: ", "Lightweight, unopinionated REST API framework"),
            ("express-validator: ", "Server-side body sanitization & validation"),
            ("Multer 1.4: ", "Multipart form-data evidence image handler"),
            ("CORS & dotenv: ", "Cross-origin control & environment configuration")
        ]),
        ("Database & Storage", [
            ("MySQL 8.0+: ", "ACID-compliant relational database engine"),
            ("mysql2 3.9: ", "Promise-based high-performance connection pool"),
            ("7 Relational Tables: ", "Fully normalized schema with foreign keys"),
            ("InnoDB Engine: ", "Row-level locking and transaction safety"),
            ("Filesystem Storage: ", "Local storage for uploaded complaint photos")
        ]),
        ("Security & Tooling", [
            ("JSON Web Tokens (JWT): ", "Stateless HS256 authorization tokens"),
            ("bcryptjs: ", "Secure password hashing with 10 salt rounds"),
            ("Git & GitHub: ", "Version control & repository management"),
            ("Visual Studio Code: ", "Primary development environment"),
            ("Cross-Browser: ", "Tested on Chrome, Edge, and Firefox")
        ])
    ]
    
    positions_s7 = [(0.8, 1.7), (6.8, 1.7), (0.8, 4.4), (6.8, 4.4)]
    
    for idx, (cat_title, cat_items) in enumerate(stacks):
        x, y = positions_s7[idx]
        add_card(s7, x, y, 5.733, 2.5)
        tb = s7.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(5.333), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cat_title
        p.font.name = 'Arial'; p.font.size = Pt(13.5); p.font.bold = True; p.font.color.rgb = NAVY
        for it_name, it_desc in cat_items:
            p_it = tf.add_paragraph()
            p_it.space_before = Pt(3)
            r1 = p_it.add_run(); r1.text = "• " + it_name; r1.font.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = DARK
            r2 = p_it.add_run(); r2.text = it_desc; r2.font.size = Pt(10.5); r2.font.color.rgb = MUTED

    # =========================================================================
    # SLIDE 8: APPLICATION SCREENS / RESULTS
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Application Screens & Implementation Results", "Authentic screenshots captured from the live CCMS application")
    
    screenshots = [
        ("docs/screenshots/01_login_page.png", "1. Role-Based Login Portal", "Authentication entry point with demo account quick-fill info."),
        ("docs/screenshots/04_new_complaint.png", "2. Complaint Lodging Form", "Dropdown selection for categories, locations, and priority levels."),
        ("docs/screenshots/07_admin_dashboard.png", "3. Administrator Control Center", "Real-time metrics, status breakdowns, and critical open hazards."),
        ("docs/screenshots/15_maintenance_tasks.png", "4. Technician Work Order Queue", "Assigned task tracking with 'Start Work' and 'Mark Resolved' actions.")
    ]
    
    positions_s8 = [(0.8, 1.7), (6.8, 1.7), (0.8, 4.4), (6.8, 4.4)]
    
    for idx, (img_path, s_title, s_caption) in enumerate(screenshots):
        x, y = positions_s8[idx]
        add_card(s8, x, y, 5.733, 2.5, fill_color=WHITE)
        if os.path.exists(img_path):
            s8.shapes.add_picture(img_path, Inches(x + 0.1), Inches(y + 0.1), width=Inches(3.3), height=Inches(2.3))
        
        tb = s8.shapes.add_textbox(Inches(x + 3.5), Inches(y + 0.3), Inches(2.1), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.name = 'Arial'; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = NAVY
        p2 = tf.add_paragraph()
        p2.text = s_caption
        p2.font.name = 'Arial'; p2.font.size = Pt(10); p2.font.color.rgb = MUTED
        p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 9: TESTING & RESULTS
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Testing & Quality Verification", "Rigorous verification covering unit logic, state progression, and end-to-end workflows")
    
    # Left Card: Testing Methodology
    add_card(s9, 0.8, 1.7, 4.5, 5.2)
    tb_tm = s9.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(4.1), Inches(4.8))
    tf_tm = tb_tm.text_frame
    tf_tm.word_wrap = True
    
    p = tf_tm.paragraphs[0]
    p.text = "Testing Methodology"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = NAVY
    
    t_methods = [
        ("Unit Testing: ", "Verified bcrypt password hashing, token decoding, and unique tracking code generators."),
        ("Basis Path Testing: ", "Mapped cyclomatic complexity of statusController.js; validated all legal and illegal transition paths."),
        ("Role Access Testing: ", "Verified that student tokens cannot access admin or technician routes (HTTP 403 Forbidden)."),
        ("System Integration: ", "Automated browser validation via Playwright confirming end-to-end user journeys.")
    ]
    for m_name, m_desc in t_methods:
        p_m = tf_tm.add_paragraph()
        p_m.space_before = Pt(8)
        r1 = p_m.add_run(); r1.text = "• " + m_name; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = BLUE
        r2 = p_m.add_run(); r2.text = m_desc; r2.font.size = Pt(10.5); r2.font.color.rgb = MUTED

    # Right Card: Test Results Table
    t_shape9 = s9.shapes.add_table(7, 4, Inches(5.6), Inches(1.7), Inches(6.933), Inches(5.2))
    tbl9 = t_shape9.table
    tbl9.columns[0].width = Inches(0.9)
    tbl9.columns[1].width = Inches(2.2)
    tbl9.columns[2].width = Inches(2.833)
    tbl9.columns[3].width = Inches(1.0)
    
    t9_headers = ["Test ID", "Scenario Evaluated", "Expected & Observed Result", "Status"]
    for c_idx, h_text in enumerate(t9_headers):
        cell = tbl9.cell(0, c_idx)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h_text; p.font.name = 'Arial'; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
    t9_rows = [
        ["TC01", "Administrator Authentication", "HTTP 200; JWT issued; redirects to Admin Dashboard", "PASS"],
        ["TC02", "Student User Authentication", "HTTP 200; JWT issued; redirects to Student Dashboard", "PASS"],
        ["TC03", "Complaint Lodging with Photo", "HTTP 201; Reference code CMP-ID generated; stored in DB", "PASS"],
        ["TC04", "Illegal Transition (NEW -> CLOSED)", "HTTP 400; Transition rejected by FSM validator", "PASS"],
        ["TC05", "Admin Technician Assignment", "HTTP 200; Complaint assigned; dispatched to staff queue", "PASS"],
        ["TC06", "Unauthorized Endpoint Access", "HTTP 403 Forbidden; Student token blocked from admin API", "PASS"]
    ]
    for r_idx, row_vals in enumerate(t9_rows):
        for c_idx, val in enumerate(row_vals):
            cell = tbl9.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 0 else LIGHT_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = 'Arial'; p.font.size = Pt(10)
            if c_idx == 0:
                p.alignment = PP_ALIGN.CENTER; p.font.bold = True
            elif c_idx == 3:
                p.alignment = PP_ALIGN.CENTER; p.font.bold = True; p.font.color.rgb = GREEN

    # =========================================================================
    # SLIDE 10: CONCLUSION & FUTURE SCOPE
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Conclusion & Future Scope", "Summary of engineering achievements and potential future enhancements")
    
    # Left Card: Conclusion
    add_card(s10, 0.8, 1.7, 5.7, 4.3)
    tb_c = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(3.9))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    
    p = tf_c.paragraphs[0]
    p.text = "Project Conclusion"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = NAVY
    
    concl_points = [
        ("Paperless Digitization: ", "Replaced error-prone manual logbooks with an efficient 24/7 web-based portal."),
        ("Accountability & Transparency: ", "Enforced role-based access and an immutable audit trail for every status change."),
        ("Structured Workflow: ", "Implemented an explicit Finite State Machine preventing illegal state jumps."),
        ("Multi-Role Empowerment: ", "Tailored dashboards for Complainants, Maintenance Technicians, and Administrators.")
    ]
    for c_title, c_text in concl_points:
        p_c = tf_c.add_paragraph()
        p_c.space_before = Pt(6)
        r1 = p_c.add_run(); r1.text = "✓ " + c_title; r1.font.bold = True; r1.font.size = Pt(11.5); r1.font.color.rgb = GREEN
        r2 = p_c.add_run(); r2.text = c_text; r2.font.size = Pt(11); r2.font.color.rgb = MUTED

    # Right Card: Future Scope
    add_card(s10, 6.8, 1.7, 5.7, 4.3)
    tb_f = s10.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(3.9))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    
    p = tf_f.paragraphs[0]
    p.text = "Future Scope"
    p.font.name = 'Arial'; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = BLUE
    
    scope_points = [
        ("Automated Notifications: ", "Integration of SMTP email alerts and SMS gateways (Twilio / Nodemailer)."),
        ("Cloud Object Storage: ", "Migration of image uploads from local filesystem to Amazon S3 or Cloudinary."),
        ("Native Mobile App: ", "Cross-platform mobile application development using React Native."),
        ("Automated Dispatching: ", "Intelligent staff allocation based on technician trade skills and active workload.")
    ]
    for s_title, s_text in scope_points:
        p_s = tf_f.add_paragraph()
        p_s.space_before = Pt(6)
        r1 = p_s.add_run(); r1.text = "• " + s_title; r1.font.bold = True; r1.font.size = Pt(11.5); r1.font.color.rgb = DARK
        r2 = p_s.add_run(); r2.text = s_text; r2.font.size = Pt(11); r2.font.color.rgb = MUTED

    # Bottom Thank You Banner
    thank_card = add_card(s10, 0.8, 6.2, 11.733, 0.9, fill_color=NAVY, border_color=NAVY)
    tb_ty = s10.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.333), Inches(0.8))
    tf_ty = tb_ty.text_frame
    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "Thank You  |  Questions & Answers"
    p_ty.font.name = 'Arial'; p_ty.font.size = Pt(18); p_ty.font.bold = True; p_ty.font.color.rgb = WHITE
    p_ty.alignment = PP_ALIGN.CENTER

    output_pptx = "Campus_Complaint_Management_System_Mini_Project_Presentation.pptx"
    prs.save(output_pptx)
    print(f"Presentation successfully generated at: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
